# -*- coding: utf-8 -*-
"""Publish the ERV cost database to the web tool, in one step.

Same shape as the dashboard's build: read the workbook that is the real
source of truth, write the artefact the page reads, push. Nobody uploads
anything to look at it - they open the link.

The difference from the dashboard's build is what goes in the artefact.
The dashboard publishes its rows as they are, which is why its data.json
answers a plain GET today with eighty client project names, design fees
and actual revenue. This one anonymises by default and encrypts by
default, so the file that lands in a public repository is prices with
nobody's name on them, and unreadable without the passphrase.

Two switches at the top change that. They are one word each, and the
consequence of flipping them is written next to them rather than left to
be discovered.

    python _src/build_cost.py            build, then ask before pushing
    python _src/build_cost.py --push     build and push without asking
    python _src/build_cost.py --no-git   build the file only
"""
import argparse, base64, getpass, hashlib, io, json, os, subprocess, sys

# openpyxl warns about the data validation extension on every read of this
# workbook. It is about a feature we do not use, it cannot be acted on, and
# it prints above the banner where it reads like an error to whoever just
# double-clicked the .bat.
import warnings
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

# the console is cp1252 until told otherwise, and every message here is Thai
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

# ══════════════════════════════════════════════════════════════════════
#  the two decisions, in one place
# ══════════════════════════════════════════════════════════════════════

ENCRYPT = True
#   True   ไฟล์ที่ push เป็น ciphertext · เปิดลิงก์แล้วหน้าเว็บถามรหัสครั้งเดียว
#          (ติ๊ก "จำรหัสในเครื่องนี้" แล้วครั้งต่อไปไม่ถามอีก)
#   False  ไฟล์เป็น JSON ธรรมดาเหมือน data.json ของ dashboard · เปิดลิงก์เห็นเลย
#          ไม่ต้องใส่อะไร  ⚠️ และใครก็ตามที่รู้ URL โหลดไปอ่านได้ทั้งก้อน

ANONYMISE = True
#   True   ตัดชื่อโปรเจค ชื่อลูกค้า จังหวัด หมายเหตุ และ README ออกก่อนเผยแพร่
#          ตัวเลขครบเหมือนเดิม RateCard กับ Estimator ทำงานเท่าเดิมทุกอย่าง
#   False  เผยแพร่ทั้งฐานรวมชื่อลูกค้า

REPO   = r'C:\tmp\dmd'
SOURCE = (r'C:\Users\coral\OneDrive - Coral Life\Coral Life - BA - Solution Design'
          r'\1_Non-Project\7_Engineer\IAQ Solution\0_ERV Management'
          r'\1_Design Management\ERV_Project_Cost_Database.xlsx')

ITERS = 310000
LOG = []


def say(s):
    LOG.append(s)


# ══════════════════════════════════════════════════════════════════════
#  reading the workbook - the same cells the web page reads
# ══════════════════════════════════════════════════════════════════════
def read_workbook(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)

    ws = wb['Projects']
    head = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    cols = [h for h in head if h]
    projects = []
    for r in range(3, ws.max_row + 1):
        row = {}
        any_value = False
        for i, name in enumerate(cols):
            v = ws.cell(row=r, column=i + 1).value
            if v not in (None, ''):
                any_value = True
            row[name] = '' if v is None else v
        if any_value and row.get(cols[0]):
            projects.append(row)

    boq = []
    if 'BOQ_Summary' in wb.sheetnames:
        bs = wb['BOQ_Summary']
        for r in range(3, bs.max_row + 1):
            pid = bs.cell(row=r, column=1).value
            if not pid:
                continue
            def cell(c):
                return bs.cell(row=r, column=c).value
            amount = cell(6)
            boq.append({
                'pid': str(pid), 'cat': float(cell(2) or 0),
                'item': '' if cell(3) is None else str(cell(3)),
                'part': '' if cell(5) is None else str(cell(5)),
                'amount': float(amount) if isinstance(amount, (int, float)) else 0.0,
                'basis': str(cell(7) or 'BOQ price'),
                'remark': '' if cell(8) is None else str(cell(8))})

    lists, cats = {}, []
    if 'Lists' in wb.sheetnames:
        ls = wb['Lists']
        for c in range(1, ls.max_column + 1):
            name = ls.cell(row=2, column=c).value
            if not name:
                continue
            vals = []
            for r in range(3, 21):
                v = ls.cell(row=r, column=c).value
                if v not in (None, ''):
                    vals.append(str(v))
            lists[str(name)] = vals
        for r in range(24, 41):
            code = ls.cell(row=r, column=1).value
            if code in (None, ''):
                continue
            cats.append({'code': float(code),
                         'name': str(ls.cell(row=r, column=2).value or ''),
                         'group': str(ls.cell(row=r, column=3).value or ''),
                         'driver': str(ls.cell(row=r, column=4).value or '')})

    readme = []
    if 'README' in wb.sheetnames:
        rs = wb['README']
        for r in range(1, min(rs.max_row, 200) + 1):
            v = rs.cell(row=r, column=2).value
            readme.append('' if v is None else str(v))
        while readme and readme[-1] == '':
            readme.pop()

    say('อ่านไฟล์แล้ว · %d โปรเจค · %d แถวหมวดต้นทุน · %d รายการ dropdown · README %d บรรทัด'
        % (len(projects), len(boq), len(lists), len(readme)))
    return {'v': 1, 'lists': lists, 'cats': cats, 'projects': projects,
            'boq': boq, 'readme': readme}


def anonymise(db):
    for i, p in enumerate(db['projects']):
        p['Project_Name'] = 'โปรเจคที่ %d' % (i + 1)
        for k in ('Client', 'Province', 'Notes'):
            if k in p:
                p[k] = ''
    db.pop('readme', None)
    db['anon'] = True
    say('ปกปิดชื่อแล้ว · ไม่มีชื่อโปรเจค ลูกค้า จังหวัด หมายเหตุ และไม่มี README ในไฟล์')
    return db


def encrypt(payload, passphrase):
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    salt, iv = os.urandom(16), os.urandom(12)
    key = hashlib.pbkdf2_hmac('sha256', passphrase.encode('utf-8'), salt, ITERS, dklen=32)
    raw = json.dumps(payload, ensure_ascii=False).encode('utf-8')
    ct = AESGCM(key).encrypt(iv, raw, None)
    say('เข้ารหัสแล้ว · AES-256-GCM · PBKDF2-SHA256 %s รอบ' % format(ITERS, ','))
    return {'format': 'coral.enc', 'v': 1, 'alg': 'AES-256-GCM',
            'kdf': {'name': 'PBKDF2', 'hash': 'SHA-256', 'iters': ITERS,
                    'salt': base64.b64encode(salt).decode()},
            'iv': base64.b64encode(iv).decode(),
            'ct': base64.b64encode(ct).decode(),
            'builtAt': __import__('datetime').datetime.now().isoformat(timespec='seconds')}


def git(args, cwd):
    p = subprocess.run(['git'] + args, cwd=cwd, capture_output=True, text=True,
                       encoding='utf-8', errors='replace')
    return p.returncode, (p.stdout or '') + (p.stderr or '')


def banner():
    print('')
    print('  ══════════════════════════════════════════════════════════════')
    print('    ERV COST DATABASE  ·  อัปเดตขึ้นเว็บ')
    print('  ══════════════════════════════════════════════════════════════')
    print('')
    print('    อ่าน   %s' % os.path.basename(SOURCE))
    print('    เขียน  cost/%s ใน repo แล้ว push'
          % ('data.enc.json' if ENCRYPT else 'data.json'))
    print('')
    print('    ปิด Excel ก่อนถ้าเปิดไฟล์นี้ค้างไว้ · ใช้เวลาประมาณหนึ่งนาที')
    print('')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--push', action='store_true', help='ไม่ต้องถามก่อน push')
    ap.add_argument('--no-git', action='store_true', help='สร้างไฟล์อย่างเดียว')
    args = ap.parse_args()

    banner()

    if not os.path.exists(SOURCE):
        say('หาไฟล์ต้นทางไม่เจอ: %s' % SOURCE)
        return 1
    if not os.path.isdir(os.path.join(REPO, '.git')):
        say('หา repo ไม่เจอที่ %s' % REPO)
        return 1

    db = read_workbook(SOURCE)
    if ANONYMISE:
        db = anonymise(db)
    else:
        say('⚠️ โหมดทั้งฐาน · ไฟล์ที่เผยแพร่มีชื่อโปรเจคและลูกค้าอยู่ด้วย')

    payload = {'format': 'coral.cost', 'v': 1,
               'mode': 'anon' if ANONYMISE else 'full',
               'savedAt': __import__('datetime').datetime.now().isoformat(timespec='seconds'),
               'db': db}

    out_dir = os.path.join(REPO, 'cost')
    enc_path = os.path.join(out_dir, 'data.enc.json')
    plain_path = os.path.join(out_dir, 'data.json')

    if ENCRYPT:
        pw = os.environ.get('CORAL_COST_PASS') or getpass.getpass('รหัสผ่านสำหรับไฟล์เผยแพร่: ')
        if len(pw) < 12:
            say('รหัสสั้นเกินไป ต้องอย่างน้อย 12 ตัวอักษร')
            return 1
        body = encrypt(payload, pw)
        target, gone = enc_path, plain_path
    else:
        say('⚠️ โหมดไม่เข้ารหัส · ใครรู้ URL ก็โหลดไฟล์นี้ไปอ่านได้ทั้งก้อน')
        body, target, gone = payload, plain_path, enc_path

    io.open(target, 'w', encoding='utf-8').write(json.dumps(body, ensure_ascii=False))
    say('เขียน %s · %s KB' % (os.path.basename(target),
                              format(os.path.getsize(target) // 1024, ',')))
    if os.path.exists(gone):
        os.remove(gone)
        say('ลบ %s ของโหมดเดิมออกแล้ว' % os.path.basename(gone))

    if args.no_git:
        return 0

    code, out = git(['pull', '--rebase', '--quiet'], REPO)
    if code != 0:
        say('git pull ไม่ผ่าน · %s' % out.strip()[:400])
        return 1
    git(['add', 'cost'], REPO)
    code, out = git(['diff', '--cached', '--quiet'], REPO)
    if code == 0:
        say('ข้อมูลไม่เปลี่ยนจากครั้งก่อน · ไม่มีอะไรต้อง push')
        return 0

    if not args.push:
        print('\n'.join(LOG))
        ans = input('\npush ขึ้นเว็บเลยไหม [y/N] ')
        if ans.strip().lower() not in ('y', 'yes'):
            git(['restore', '--staged', 'cost'], REPO)
            say('ยกเลิกแล้ว · ไฟล์ยังอยู่ในเครื่อง ไม่ได้ push')
            return 0

    code, out = git(['commit', '-m', 'Publish cost database'], REPO)
    if code != 0:
        say('commit ไม่ผ่าน · %s' % out.strip()[:400])
        return 1
    code, out = git(['push', '--quiet'], REPO)
    if code != 0:
        say('push ไม่ผ่าน · %s' % out.strip()[:400])
        return 1
    say('push แล้ว · เว็บจะอัปเดตในประมาณหนึ่งนาที')
    say('https://nchatupongtrakul-ops.github.io/design-management-dashboard/cost/')
    return 0


if __name__ == '__main__':
    rc = 1
    try:
        rc = main()
    except Exception as e:                       # noqa: BLE001 - shown to the user
        say('ผิดพลาด · %s' % e)
    io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'build_cost.log'),
            'w', encoding='utf-8').write('\n'.join(LOG))
    try:
        print('\n'.join(LOG))
    except Exception:
        print('done - see _src/build_cost.log')
    sys.exit(rc)
